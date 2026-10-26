"""Read-only parsers for user-supplied request documents."""

from __future__ import annotations

import asyncio
import csv
import mimetypes
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Protocol

from docx import Document
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from pydantic import Field
from pypdf import PdfReader

from .exceptions import DocumentParseError
from .schemas import (
    ImageOCRDTO,
    ParsedDocumentDTO,
    ParsedTableDTO,
    StrictDTO,
    StructuredLLMResultDTO,
)


class DocumentParserConfig(StrictDTO):
    max_file_size_bytes: int = Field(default=50 * 1024 * 1024, ge=1)
    max_cells: int = Field(default=100_000, ge=1)
    max_text_characters: int = Field(default=2_000_000, ge=1)
    max_image_pixels: int = Field(default=40_000_000, ge=1)
    max_pdf_ocr_pages: int = Field(default=30, ge=1, le=200)
    pdf_ocr_dpi: int = Field(default=240, ge=72, le=300)


class ImageTextExtractor(Protocol):
    async def extract_image_text(
        self, image_bytes: bytes, mime_type: str
    ) -> StructuredLLMResultDTO[ImageOCRDTO]: ...


class DocumentParserService:
    SUPPORTED_IMAGES = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}

    def __init__(
        self,
        *,
        image_text_extractor: ImageTextExtractor | None = None,
        config: DocumentParserConfig | None = None,
    ) -> None:
        self._image_text_extractor = image_text_extractor
        self.config = config or DocumentParserConfig()

    async def parse(self, path: Path) -> ParsedDocumentDTO:
        resolved = path.resolve()
        if not resolved.is_file():
            raise DocumentParseError(f"Document not found: {resolved}")
        size = resolved.stat().st_size
        if size > self.config.max_file_size_bytes:
            raise DocumentParseError(
                f"Document exceeds {self.config.max_file_size_bytes} bytes"
            )

        suffix = resolved.suffix.lower()
        try:
            if suffix == ".docx":
                parsed = self._parse_docx(resolved)
            elif suffix == ".pdf":
                parsed = await self._parse_pdf(resolved)
            elif suffix in {".xlsx", ".xlsm"}:
                parsed = self._parse_xlsx(resolved)
            elif suffix in {".txt", ".csv"}:
                parsed = self._parse_text(resolved)
            elif suffix in self.SUPPORTED_IMAGES:
                parsed = await self._parse_image(resolved)
            elif suffix in {".doc", ".xls"}:
                parsed = await self._parse_legacy_office(resolved)
            else:
                raise DocumentParseError(f"Unsupported document type: {suffix}")
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError(
                f"Could not parse {resolved.name}: {exc}"
            ) from exc

        if len(parsed.text) > self.config.max_text_characters:
            parsed.text = parsed.text[: self.config.max_text_characters]
            parsed.warnings.append("text_truncated")
        return parsed

    async def _parse_legacy_office(self, path: Path) -> ParsedDocumentDTO:
        """Convert legacy binary Office files before parsing their contents."""

        target_format = "docx" if path.suffix.lower() == ".doc" else "xlsx"
        with tempfile.TemporaryDirectory(prefix="buscript-office-") as temp:
            profile = Path(temp) / "libreoffice-profile"
            profile.mkdir()
            environment = os.environ.copy()
            environment["HOME"] = temp
            process = await asyncio.create_subprocess_exec(
                "soffice",
                "--headless",
                f"-env:UserInstallation={profile.as_uri()}",
                "--convert-to",
                target_format,
                "--outdir",
                temp,
                str(path),
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.PIPE,
                env=environment,
            )
            try:
                stdout, stderr = await asyncio.wait_for(
                    process.communicate(), timeout=60
                )
            except TimeoutError as exc:
                process.kill()
                await process.communicate()
                raise DocumentParseError(
                    f"Legacy {path.suffix.lower()} conversion timed out"
                ) from exc
            converted = Path(temp) / f"{path.stem}.{target_format}"
            if process.returncode != 0 or not converted.is_file():
                detail = (stderr or stdout).decode("utf-8", errors="replace").strip()
                raise DocumentParseError(
                    f"Could not convert legacy {path.suffix.lower()}: "
                    f"{detail or 'LibreOffice did not create a file'}"
                )
            parsed = (
                self._parse_docx(converted)
                if target_format == "docx"
                else self._parse_xlsx(converted)
            )
            parsed.path = path
            parsed.warnings.append(f"legacy_{path.suffix.lower()[1:]}_converted")
            return parsed

    def _parse_docx(self, path: Path) -> ParsedDocumentDTO:
        document = Document(path)
        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        tables: list[ParsedTableDTO] = []
        table_texts: list[str] = []
        cell_count = 0
        for table_index, table in enumerate(document.tables, start=1):
            rows: list[list[str]] = []
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                cell_count += len(values)
                if cell_count > self.config.max_cells:
                    raise DocumentParseError("DOCX table cell limit exceeded")
                rows.append(values)
                table_texts.append("\t".join(values))
            tables.append(
                ParsedTableDTO(source=f"table:{table_index}", rows=rows)
            )
        text = "\n".join([*paragraphs, *table_texts]).strip()
        return ParsedDocumentDTO(
            path=path,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            text=text,
            tables=tables,
        )

    async def _parse_pdf(self, path: Path) -> ParsedDocumentDTO:
        if not path.read_bytes()[:5] == b"%PDF-":
            raise DocumentParseError("File extension is PDF but signature is not")
        reader = PdfReader(path, strict=False)
        page_texts: list[str] = []
        empty_pages = 0
        for page in reader.pages:
            text = (page.extract_text() or "").strip()
            if not text:
                empty_pages += 1
            page_texts.append(text)
        warnings: list[str] = []
        if empty_pages:
            warnings.append(f"pages_without_text_layer:{empty_pages}")
        if page_texts and empty_pages == len(page_texts):
            warnings.append("scanned_pdf_ocr_required")
        provider = None
        model = None
        confidence = None
        if empty_pages and self._image_text_extractor is not None:
            (
                ocr_texts,
                ocr_confidences,
                provider,
                model,
                ocr_warnings,
            ) = await self._ocr_pdf_pages(path, page_texts)
            warnings.extend(ocr_warnings)
            for index, text in ocr_texts.items():
                page_texts[index] = text
            if ocr_texts:
                warnings = [
                    warning
                    for warning in warnings
                    if warning != "scanned_pdf_ocr_required"
                ]
                warnings.append("scanned_pdf_ocr_applied")
            if ocr_confidences:
                confidence = sum(ocr_confidences) / len(ocr_confidences)
        return ParsedDocumentDTO(
            path=path,
            media_type="application/pdf",
            text=(
                "\n\n".join(
                    f"--- Страница {index} ---\n{text}"
                    for index, text in enumerate(page_texts, start=1)
                ).strip()
                if len(page_texts) > 1
                else "\n\n".join(page_texts).strip()
            ),
            page_count=len(reader.pages),
            confidence=confidence,
            provider=provider,
            model=model,
            warnings=warnings,
        )

    async def _ocr_pdf_pages(
        self,
        path: Path,
        page_texts: list[str],
    ) -> tuple[
        dict[int, str],
        list[float],
        str | None,
        str | None,
        list[str],
    ]:
        pdftoppm = shutil.which("pdftoppm")
        if pdftoppm is None:
            return {}, [], None, None, ["pdf_ocr_renderer_unavailable"]
        empty_indexes = [
            index for index, text in enumerate(page_texts) if not text.strip()
        ]
        if len(empty_indexes) > self.config.max_pdf_ocr_pages:
            empty_indexes = empty_indexes[: self.config.max_pdf_ocr_pages]
            warnings = ["pdf_ocr_page_limit_reached"]
        else:
            warnings = []
        output: dict[int, str] = {}
        confidences: list[float] = []
        provider = None
        model = None
        with tempfile.TemporaryDirectory(prefix="buscript-pdf-ocr-") as temp:
            prefix = Path(temp) / "page"
            command = [
                pdftoppm,
                "-png",
                "-r",
                str(self.config.pdf_ocr_dpi),
                str(path),
                str(prefix),
            ]
            result = await asyncio.to_thread(
                subprocess.run,
                command,
                capture_output=True,
                check=False,
                timeout=120,
            )
            if result.returncode != 0:
                return {}, [], None, None, [
                    *warnings,
                    "pdf_ocr_render_failed",
                ]
            rendered = sorted(Path(temp).glob("page-*.png"))
            for index in empty_indexes:
                if index >= len(rendered):
                    warnings.append(f"pdf_ocr_page_missing:{index + 1}")
                    continue
                image_path = rendered[index]
                try:
                    ocr = await self._image_text_extractor.extract_image_text(
                        image_path.read_bytes(),
                        "image/png",
                    )
                except Exception:
                    warnings.append(f"pdf_ocr_page_failed:{index + 1}")
                    continue
                if ocr.data.text.strip():
                    output[index] = ocr.data.text.strip()
                    confidences.append(ocr.data.confidence)
                    provider = provider or ocr.provider
                    model = model or ocr.actual_model
        return output, confidences, provider, model, warnings

    def _parse_xlsx(self, path: Path) -> ParsedDocumentDTO:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        tables: list[ParsedTableDTO] = []
        lines: list[str] = []
        cell_count = 0
        try:
            for sheet in workbook.worksheets:
                rows: list[list[str]] = []
                lines.append(f"=== Лист: {sheet.title} ===")
                for raw_row in sheet.iter_rows(values_only=True):
                    values = [
                        "" if value is None else str(value).strip()
                        for value in raw_row
                    ]
                    while values and values[-1] == "":
                        values.pop()
                    if not values:
                        continue
                    cell_count += len(values)
                    if cell_count > self.config.max_cells:
                        raise DocumentParseError("XLSX cell limit exceeded")
                    rows.append(values)
                    lines.append("\t".join(values))
                tables.append(ParsedTableDTO(source=sheet.title, rows=rows))
        finally:
            workbook.close()
        media_type = (
            "application/vnd.ms-excel.sheet.macroEnabled.12"
            if path.suffix.lower() == ".xlsm"
            else "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
        return ParsedDocumentDTO(
            path=path,
            media_type=media_type,
            text="\n".join(lines).strip(),
            tables=tables,
            sheet_count=len(tables),
        )

    def _parse_text(self, path: Path) -> ParsedDocumentDTO:
        raw = path.read_bytes()
        text = ""
        for encoding in ("utf-8-sig", "cp1251", "utf-16"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            raise DocumentParseError("Could not decode text document")
        media_type = (
            "text/csv" if path.suffix.lower() == ".csv" else "text/plain"
        )
        tables: list[ParsedTableDTO] = []
        if media_type == "text/csv":
            rows = [
                [cell.strip() for cell in row]
                for row in csv.reader(text.splitlines())
                if any(cell.strip() for cell in row)
            ]
            tables.append(ParsedTableDTO(source=path.name, rows=rows))
        return ParsedDocumentDTO(
            path=path,
            media_type=media_type,
            text=text.strip(),
            tables=tables,
        )

    async def _parse_image(self, path: Path) -> ParsedDocumentDTO:
        try:
            with Image.open(path) as image:
                width, height = image.size
                image_format = (image.format or "").upper()
                if width * height > self.config.max_image_pixels:
                    raise DocumentParseError("Image pixel limit exceeded")
                image.verify()
        except (UnidentifiedImageError, OSError) as exc:
            raise DocumentParseError("Invalid or corrupted image") from exc

        media_type = Image.MIME.get(image_format) or mimetypes.guess_type(path)[0]
        media_type = media_type or "application/octet-stream"
        if not media_type.startswith("image/"):
            raise DocumentParseError("Could not determine image MIME type")

        warnings: list[str] = []
        text = ""
        confidence = None
        provider = None
        model = None
        if self._image_text_extractor is None:
            warnings.append("image_ocr_not_configured")
        else:
            result = await self._image_text_extractor.extract_image_text(
                path.read_bytes(), media_type
            )
            text = result.data.text
            confidence = result.data.confidence
            provider = result.provider
            model = result.actual_model

        return ParsedDocumentDTO(
            path=path,
            media_type=media_type,
            text=text,
            image_width=width,
            image_height=height,
            confidence=confidence,
            provider=provider,
            model=model,
            warnings=warnings,
        )
