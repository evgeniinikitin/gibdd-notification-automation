"""Build a synthetic DOCX layout for this public portfolio copy.

No customer material or source document is embedded in this template.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document


def ensure_demo_template(path: Path) -> Path:
    if path.exists():
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = "Demo notification template"
    document.add_heading("Уведомление · демонстрационный шаблон", 0)
    document.add_paragraph("Дата поездки: {{TRIP_DATE_TEXT}}")
    document.add_paragraph("Организатор: {{ORGANIZER_NAME}}")

    def table(headers: list[str], markers: list[str]) -> None:
        grid = document.add_table(rows=2, cols=len(headers))
        grid.style = "Table Grid"
        for cell, value in zip(grid.rows[0].cells, headers, strict=True):
            cell.text = value
        for cell, value in zip(grid.rows[1].cells, markers, strict=True):
            cell.text = value

    table(
        ["№", "Ребёнок", "Возраст", "Дата рождения", "Родитель"],
        ["№", "{{CHILD_FULL_NAME}}", "{{CHILD_AGE}}", "{{CHILD_BIRTH_DATE}}", "{{PARENT_PHONE}}"],
    )
    table(
        ["№", "Сопровождающий", "Роль", "Телефон"],
        ["№", "{{COMPANION_NAME}}", "{{COMPANION_ROLE}}", "{{COMPANION_PHONE}}"],
    )
    table(
        ["Автобус", "Номер", "Диагностическая карта"],
        ["{{BUS_MODEL}}", "{{BUS_PLATE}}", "{{BUS_DIAGNOSTIC_CARD}}"],
    )
    table(
        ["Водитель", "Дата рождения", "Удостоверение", "Выдано", "Стаж"],
        ["{{DRIVER_NAME}}", "{{DRIVER_BIRTH_DATE}}", "{{DRIVER_LICENSE_NUMBER}}", "{{DRIVER_LICENSE_ISSUE_DATE}}", "{{DRIVER_EXPERIENCE}}"],
    )
    table(["Маршрут"], ["{{ROUTE_DESCRIPTION_MANUAL}}"])
    document.add_paragraph("{{SIGNATORY_POSITION}} / {{SIGNATORY_NAME}} /")
    document.save(path)
    return path
